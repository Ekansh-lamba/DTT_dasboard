"""The study report, in the Apollo Tyres deck style.

Slides are built on ``templates/apollo_report_template.pptx`` -- the masters
and layouts of the DTT project charter deck, with its slides removed -- so the
report carries the same look as the charter: white 16:9 slides with the
Apollo logo, a black heading over an orange sub-heading and a grey rule,
purple-headed tables, the Apollo/Vredestein title slide and the purple-spiral
section dividers. Every slide is still built from this study's own channels
and figures; nothing from the charter's content is reused.

If the template is missing (a stripped install), a plain 16:9 deck with the
same layout geometry is produced instead, so a report is never lost to a
styling file.
"""

import datetime
import logging
from pathlib import Path
from typing import Dict, List, Optional

import matplotlib
matplotlib.use("Agg")
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt, Emu, Inches

from dtt.config import (
    MANDATORY_CHANNELS,
    WHEEL_GROUPS,
    PERCENTILES,
    RunConfig,
)
from dtt.validation.validator import ValidationReport
from dtt.channel_names import display_name
from dtt.analysis.sessions_plot import session_info

logger = logging.getLogger(__name__)

TEMPLATE = Path(__file__).resolve().parent / "templates" / "apollo_report_template.pptx"

# Layout names in the template (the charter deck's own).
_LAYOUT_CONTENT = "Heading + text x2 columns (Apollo) A"
_LAYOUT_DIVIDER = "Divider opt 2 (Apollo) A"
_LAYOUT_COVER   = "1_Title Slide opt 1 (Apolo Tyres) A"
_LAYOUT_RESULT  = "Heading + text + 3 imgs (Apollo) A"   # charter slides 9-14: "APOLLO TYRES LTD" footer
_LAYOUT_SECTION = "Title Only"                           # charter slides 8, 11, 13

# Apollo palette, from the charter deck's theme.
_PURPLE    = RGBColor(0x95, 0x19, 0xEF)     # accent1: table headers, labels
_PURPLE_DK = RGBColor(0x5C, 0x2C, 0x91)     # accent4: panel outlines, footer
_ORANGE    = RGBColor(0xED, 0x7D, 0x31)     # sub-headings
_INK       = RGBColor(0x18, 0x18, 0x1E)     # dk1: body text
_GREY      = RGBColor(0x59, 0x59, 0x59)     # notes, slide numbers
_RULE      = RGBColor(0xBF, 0xBF, 0xBF)     # heading rule
_ROW_ALT   = RGBColor(0xF3, 0xEA, 0xFD)     # alternate table rows
_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
_GREEN_RGB = RGBColor(0x1E, 0x84, 0x49)
_RED_RGB   = RGBColor(0xC0, 0x39, 0x2B)

# Result slides, from the charter deck's own shapes (slides 9-14).
_R_TITLE   = RGBColor(0x6C, 0x2D, 0xC7)     # PV titles, takeaway text
_R_PILL    = RGBColor(0xE8, 0x79, 0x2D)     # route pill
_R_NAVY    = RGBColor(0x00, 0x00, 0x99)     # section pills
_R_CYAN    = RGBColor(0x00, 0xB0, 0xF0)     # CV row label, candidate
_R_BLUE    = RGBColor(0x44, 0x72, 0xC4)     # CV row label, reference
_R_MAGENTA = RGBColor(0xBC, 0x32, 0xBE)     # heatmap subtitle, summary border
_R_BOX     = RGBColor(0xF2, 0xF2, 0xF2)     # takeaway / summary fill
_R_DIVIDE  = RGBColor(0x5B, 0x9B, 0xD5)     # heatmap halves divider
_BLACK     = RGBColor(0x00, 0x00, 0x00)
_R_FONT    = "Arial"
_R_FONT_CV = "Verdana"

# Slide geometry (16:9, 13.333 x 7.5 in). The Apollo logo sits bottom-left
# (0.1..2.63 in, from 6.37 in down), so content stops above it.
SLIDE_W = Emu(12192000)          # exact 16:9, as the template
SLIDE_H = Emu(6858000)
L = Inches(0.45)
T = Inches(1.45)
IMG_W = Inches(12.43)
IMG_H = Inches(4.85)

_FONT = "Calibri"


# ---------------------------------------------------------------- primitives

def _add_textbox(slide, text: str, left, top, width, height,
                 font_size: int = 14, bold: bool = False,
                 color: RGBColor = _INK, align=PP_ALIGN.LEFT,
                 italic: bool = False):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = _FONT
        run.font.color.rgb = color
    return txb


def _layout(prs, name: str):
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    # plain fallback: the emptiest layout there is
    return min((l for m in prs.slide_masters for l in m.slide_layouts),
               key=lambda l: len(l.placeholders))


def _drop_placeholders(slide, keep=()) -> None:
    """Remove placeholders we do not fill -- an empty one shows "Click to add
    text" to whoever opens the deck."""
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in keep:
            ph._element.getparent().remove(ph._element)


def _new_presentation():
    if TEMPLATE.exists():
        prs = Presentation(str(TEMPLATE))
    else:
        logger.warning("Report template missing (%s) -- using a plain deck", TEMPLATE)
        prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


class _Deck:
    """Adds slides in the charter's three styles and numbers them."""

    def __init__(self, prs, footer: str):
        self.prs = prs
        self.footer = footer

    def _number(self, slide) -> None:
        n = len(self.prs.slides)
        _add_textbox(slide, str(n), Inches(12.35), Inches(6.9), Inches(0.6),
                     Inches(0.3), font_size=10, color=_GREY, align=PP_ALIGN.RIGHT)

    def content(self, title: str, subtitle: str = ""):
        """Black heading, orange sub-heading, grey rule -- charter slides 3-6.
        The heading is the slide's first text shape, so it reads as its title."""
        slide = self.prs.slides.add_slide(_layout(self.prs, _LAYOUT_CONTENT))
        _drop_placeholders(slide)
        _add_textbox(slide, title, L, Inches(0.22), IMG_W, Inches(0.6),
                     font_size=28, bold=True, color=_INK)
        if subtitle:
            _add_textbox(slide, subtitle, L, Inches(0.8), IMG_W, Inches(0.4),
                         font_size=16, bold=True, color=_ORANGE)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, Inches(1.25),
                                      IMG_W, Inches(0.03))
        rule.fill.solid()
        rule.fill.fore_color.rgb = _RULE
        rule.line.fill.background()
        if self.footer:
            _add_textbox(slide, self.footer, Inches(2.75), Inches(6.9),
                         Inches(8.0), Inches(0.35), font_size=11, italic=True,
                         color=_PURPLE_DK)
        self._number(slide)
        return slide

    def divider(self, title: str, subtitle: str = ""):
        """Purple-spiral section divider."""
        slide = self.prs.slides.add_slide(_layout(self.prs, _LAYOUT_DIVIDER))
        placed = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        if 0 in placed:
            placed[0].text = title
        else:
            _add_textbox(slide, title, L, Inches(2.1), Inches(6), Inches(0.7),
                         font_size=32, bold=True, color=_PURPLE)
        if 10 in placed and subtitle:
            placed[10].text = subtitle
        elif subtitle:
            _add_textbox(slide, subtitle, L, Inches(2.8), Inches(6), Inches(0.6),
                         font_size=18, color=_PURPLE_DK)
        _drop_placeholders(slide, keep={0} | ({10} if subtitle else set()))
        self._number(slide)
        return slide

    def section(self, title: str):
        """Plain centred-title section slide -- charter slides 8 and 11."""
        slide = self.prs.slides.add_slide(_layout(self.prs, _LAYOUT_SECTION))
        placed = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        if 0 in placed:
            placed[0].text = title
        else:
            _add_textbox(slide, title, Inches(0.92), Inches(0.4), Inches(11.5),
                         Inches(1.45), font_size=40)
        _drop_placeholders(slide, keep={0})
        return slide

    def result(self):
        """An empty result slide on the charter's result layout, whose
        "APOLLO TYRES LTD" footer comes with it. Its title is the first shape
        the caller adds."""
        slide = self.prs.slides.add_slide(_layout(self.prs, _LAYOUT_RESULT))
        _drop_placeholders(slide)
        return slide

    def cover(self, title: str, lines: List[str]):
        """Apollo / Vredestein title slide."""
        slide = self.prs.slides.add_slide(_layout(self.prs, _LAYOUT_COVER))
        placed = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        if 0 in placed:
            placed[0].text = title
        else:
            _add_textbox(slide, title, Inches(2.2), Inches(2.9), Inches(9), Inches(1.6),
                         font_size=36, bold=True, color=_PURPLE, align=PP_ALIGN.CENTER)
        _drop_placeholders(slide, keep={0})
        _add_textbox(slide, "\n".join(lines), Inches(2.2), Inches(4.55), Inches(9),
                     Inches(1.4), font_size=16, color=_PURPLE_DK, align=PP_ALIGN.CENTER)
        return slide


def _add_picture_fit(slide, img_path: Path, left, top, width, height):
    """Place the picture as large as fits the box, centred, aspect kept.

    Giving python-pptx both a width and a height stretches the image to the
    box: a tall box plot squeezed into a wide slot, text and all.
    """
    from PIL import Image
    with Image.open(img_path) as im:
        iw, ih = im.size
    scale = min(int(width) / iw, int(height) / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(img_path), Emu(int(left) + (int(width) - w) // 2),
                             Emu(int(top) + (int(height) - h) // 2), Emu(w), Emu(h))


def _insert_image_or_message(slide, img_path: Path, left, top, width, height,
                             reason: str = "") -> bool:
    """Insert ``img_path`` if it exists, is non-empty, and decodes as an
    image; otherwise leave an explanatory textbox in its place instead of a
    bare title with nothing under it. Returns whether the image was inserted."""
    if img_path.exists() and img_path.stat().st_size > 0:
        try:
            _add_picture_fit(slide, img_path, left, top, width, height)
            return True
        except Exception as exc:
            logger.warning("Image could not be opened, skipping: %s (%s)", img_path.name, exc)
            msg = f"Not available: {img_path.name} could not be opened ({exc})."
    else:
        logger.warning("Image not found, skipping: %s", img_path.name)
        msg = reason or f"Not available: {img_path.name} was not generated for this study."
    _add_textbox(slide, msg, left, top, width, min(height, Emu(500000)),
                 font_size=12, italic=True, color=_GREY)
    return False


# --------------------------------------------------------------------- tables

def _style_cell(cell, text: str, *, header=False, label=False, alt=False,
                size=10, color: Optional[RGBColor] = None, bold=False,
                align=PP_ALIGN.LEFT) -> None:
    cell.text = str(text)
    cell.fill.solid()
    cell.fill.fore_color.rgb = (_PURPLE if (header or label)
                                else (_ROW_ALT if alt else _WHITE))
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    for para in cell.text_frame.paragraphs:
        para.alignment = align
        for run in para.runs:
            run.font.name = _FONT
            run.font.size = Pt(size)
            run.font.bold = header or label or bold
            run.font.color.rgb = _WHITE if (header or label) else (color or _INK)


def _add_stats_table(slide, stats: Dict, channels: List[str], left, top, width, height) -> None:
    """``channels`` is this study's own mandatory-channel list
    (``config.run_channels.mandatory_channels``, generalized axle names like
    ``FR_Fx_2``) -- never the legacy ``FL_Fx``-style constant, which will
    never match a real study's stats keys."""
    channels = [ch for ch in channels if ch in stats]
    if not channels:
        _add_textbox(slide, "Not available: no statistics could be matched to "
                            "this study's channels.", left, top, width,
                     Emu(400000), font_size=13, italic=True, color=_GREY)
        return
    cols = ["Channel", "Count", "Mean", "Median", "Std", "Min", "Max"] + [f"P{p}" for p in PERCENTILES]
    row_h = min(int(height) // (len(channels) + 1), int(Inches(0.45)))
    table = slide.shapes.add_table(len(channels) + 1, len(cols), left, top, width,
                                   Emu(row_h * (len(channels) + 1))).table
    for ci, hdr in enumerate(cols):
        _style_cell(table.cell(0, ci), hdr, header=True, size=13,
                    align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)

    for ri, ch in enumerate(channels, start=1):
        s = stats[ch]
        row_vals = [display_name(ch), s.get("count"), s.get("mean"), s.get("median"), s.get("std"),
                    s.get("min"), s.get("max")] + [s.get(f"P{p}") for p in PERCENTILES]
        for ci, val in enumerate(row_vals):
            if ci == 0:
                text = str(val)
            elif val is None:
                text = "N/A"
            elif ci == 1:
                text = f"{int(val):,}"
            else:
                text = f"{val:.2f}"
            _style_cell(table.cell(ri, ci), text, alt=ri % 2 == 0, bold=ci == 0, size=12,
                        align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)


def _add_validation_table(slide, vr: ValidationReport, left, top, width, height) -> None:
    """Channel | Status, wrapped into side-by-side column pairs so a 40-channel
    recording still fits the slide instead of running off its bottom."""
    items = list(vr.channel_status.items())
    if not items:
        _add_textbox(slide, "No channel status was recorded.", left, top, width,
                     Emu(400000), font_size=13, italic=True, color=_GREY)
        return
    per_col = 14
    groups = -(-len(items) // per_col)
    rows = min(per_col, len(items)) + 1
    row_h = min(int(height) // rows, int(Inches(0.34)))
    table = slide.shapes.add_table(rows, 2 * groups, left, top, width,
                                   Emu(row_h * rows)).table
    for g in range(groups):
        _style_cell(table.cell(0, 2 * g), "Channel", header=True, size=12)
        _style_cell(table.cell(0, 2 * g + 1), "Status", header=True, size=12)
        for r in range(per_col):
            i = g * per_col + r
            if r + 1 >= rows:
                break
            if i >= len(items):
                _style_cell(table.cell(r + 1, 2 * g), "", alt=r % 2 == 1)
                _style_cell(table.cell(r + 1, 2 * g + 1), "", alt=r % 2 == 1)
                continue
            ch, st = items[i]
            ok = st == "PRESENT"
            _style_cell(table.cell(r + 1, 2 * g), display_name(ch), alt=r % 2 == 1, size=11)
            _style_cell(table.cell(r + 1, 2 * g + 1), f"{'✓' if ok else '✗'}  {st}",
                        alt=r % 2 == 1, size=11, bold=True,
                        color=_GREEN_RGB if ok else _RED_RGB)


def _add_key_value_table(slide, rows: List[tuple], left, top, width, height) -> None:
    """Purple label column, white value column -- charter slide 1."""
    row_h = min(int(height) // max(1, len(rows)), int(Inches(0.5)))
    table = slide.shapes.add_table(len(rows), 2, left, top, width,
                                   Emu(row_h * len(rows))).table
    table.first_row = False
    table.columns[0].width = Inches(2.3)
    table.columns[1].width = Emu(int(width) - int(Inches(2.3)))
    for r, (k, v) in enumerate(rows):
        _style_cell(table.cell(r, 0), k, label=True, size=14)
        _style_cell(table.cell(r, 1), v, size=13)


def _add_panel(slide, heading: str, lines: List[str], left, top, width, height) -> None:
    """Outlined panel with a purple heading -- charter slide 2's result boxes."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.background()
    box.line.color.rgb = _PURPLE_DK
    box.line.width = Pt(1.25)
    _add_textbox(slide, heading, left + Inches(0.15), top + Inches(0.08),
                 width - Inches(0.3), Inches(0.4), font_size=18, bold=True,
                 color=_PURPLE_DK)
    _add_textbox(slide, "\n".join(f"•  {l}" for l in lines),
                 left + Inches(0.2), top + Inches(0.5), width - Inches(0.4),
                 height - Inches(0.6), font_size=15, color=_INK)


def _grid_positions(n: int, left, top, width, height, max_cols: int = 2):
    """``n`` evenly-sized (left, top, width, height) cells, up to
    ``max_cols`` per row. Generic replacement for a hardcoded 2x2 quad --
    a fixed 4-slot grid silently drops any wheel past the 4th on a 6-wheel
    truck, since zip() truncates to the shorter of the two sequences."""
    if n <= 0:
        return []
    cols = min(max_cols, n)
    rows = -(-n // cols)  # ceil division
    cw, ch = width // cols, height // rows
    return [(left + (i % cols) * cw, top + (i // cols) * ch, cw, ch) for i in range(n)]


def _audit_slides(prs: Presentation) -> List[str]:
    """Slides that ended up with only a title and nothing else -- should be
    impossible now that every image slot falls back to an explanatory
    textbox via :func:`_insert_image_or_message`, so this is a regression
    guard, not expected to ever fire."""
    warnings = []
    for i, slide in enumerate(prs.slides, start=1):
        if slide.slide_layout.name == _LAYOUT_SECTION:
            continue                      # a section title is all a section slide holds
        n_pics   = sum(1 for sh in slide.shapes if sh.shape_type == 13)
        n_tables = sum(1 for sh in slide.shapes if sh.has_table)
        n_text   = sum(1 for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        if n_pics == 0 and n_tables == 0 and n_text <= 1:
            warnings.append(f"slide {i} has only a title, no content and no explanation")
    return warnings


# ------------------------------------------------------------- result slides
#
# The charter deck's result format (slides 8-14). Every title, heading, wheel
# label, pill, divider and summary box is a native shape an analyst can edit;
# only the plots are pictures (dtt.reporting.result_figures draws them). The
# summary boxes hold placeholder text -- the observations are the analyst's.

_OBSERVATION = "Key observation: …"


def _runs_box(slide, paragraphs, left, top, width, height, *, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, fill: Optional[RGBColor] = None, inset: bool = False):
    """Text box from ``[[(text, size, bold, color, font), ...], ...]`` --
    one inner list per paragraph, one tuple per run."""
    txb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if not inset:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if fill is not None:
        txb.fill.solid()
        txb.fill.fore_color.rgb = fill
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, color, font in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = font
            r.font.color.rgb = color
    return txb


def _label(slide, text, left, top, width=1.0, height=0.3, size=14, color=_BLACK,
           align=PP_ALIGN.LEFT, font=_R_FONT):
    return _runs_box(slide, [[(text, size, True, color, font)]], left, top, width, height,
                     align=align)


def _pill(slide, text, left, top, width, height, fill: RGBColor, size: int):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                  Inches(width), Inches(height))
    pill.fill.solid()
    pill.fill.fore_color.rgb = fill
    pill.line.fill.background()
    tf = pill.text_frame
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = _R_FONT
    r.font.color.rgb = _WHITE
    return pill


def _rule(slide, x1, y1, x2, y2, color=_BLACK, width_pt=1.25):
    from pptx.enum.shapes import MSO_CONNECTOR
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width_pt)
    return line


def _note_box(slide, paragraphs, left, top, width, height):
    """Light-grey takeaway box in bold purple -- placeholder text only."""
    return _runs_box(slide, [[(t, 10.5, True, _R_TITLE, _R_FONT)] for t in paragraphs],
                     left, top, width, height, fill=_R_BOX, inset=True,
                     anchor=MSO_ANCHOR.MIDDLE)


def _result_plot(slide, fig: Path, left, top, size_in, label: str, missing: str) -> None:
    """A plot picture at its own size, with the wheel/side label written over
    its top-left corner as an editable text box, as the charter deck does."""
    placed = _insert_image_or_message(slide, fig, Inches(left), Inches(top),
                                      Inches(size_in[0]), Inches(size_in[1]), reason=missing)
    if placed and label:
        _label(slide, label, left + 0.55, top + 0.12, 0.9, 0.3, size=14)


def _vs(spec) -> str:
    return f"{spec.label_a} vs {spec.label_b}" if spec.compared else spec.label_a


def _pv_title(slide, spec, what: str) -> None:
    _label(slide, f"Force Distribution - {_vs(spec)} – {what}", 0.13, 0.02, 11.6, 0.5,
           size=28, color=_R_TITLE)
    if spec.route:
        _pill(slide, spec.route, 11.87, 0.23, 1.35, 0.34, _R_PILL, 18)


def _missing(wheel: str, comp: str) -> str:
    return f"{wheel} {comp}: not recorded in this study."


def _pv_page_fy_fz(deck: "_Deck", spec, config) -> None:
    from dtt.reporting.result_figures import PV_SLOT_IN, PV_WHEELS, figure_path
    slide = deck.result()
    _pv_title(slide, spec, "Fy & Fz")
    _rule(slide, 6.74, 0.55, 6.74, 6.86)
    halves = (("Fy", 0.1, ((0.1, 1.25), (3.45, 1.25), (0.1, 4.1), (3.45, 4.1))),
              ("Fz", 6.8, ((6.8, 1.25), (10.05, 1.25), (6.8, 3.95), (10.05, 3.95))))
    for comp, x0, slots in halves:
        _label(slide, comp, x0, 0.72, 6.5, 0.3, size=16, align=PP_ALIGN.CENTER)
        for wheel, (x, y) in zip(PV_WHEELS, slots):
            _result_plot(slide, figure_path(config, f"pv_{comp}_{wheel}"), x, y,
                         PV_SLOT_IN, wheel, _missing(wheel, comp))
    _note_box(slide, [_OBSERVATION], 6.88, 6.66, 6.36, 0.52)


def _pv_page_fx_my(deck: "_Deck", spec, config) -> None:
    from dtt.reporting.result_figures import PV_SLOT_IN, PV_WHEELS, drive_wheels, figure_path
    slide = deck.result()
    _pv_title(slide, spec, "Fx")
    _rule(slide, 6.74, 0.55, 6.74, 6.86)

    _pill(slide, "By Wheel Position", 0.13, 0.65, 2.6, 0.31, _R_NAVY, 16)
    _label(slide, "Fx", 2.75, 0.62, 1.2, 0.3, size=16, align=PP_ALIGN.CENTER)
    for wheel, (x, y) in zip(PV_WHEELS, ((0.05, 1.0), (3.4, 1.0), (0.05, 4.05), (3.4, 4.05))):
        _result_plot(slide, figure_path(config, f"pv_Fx_{wheel}"), x, y, PV_SLOT_IN,
                     wheel, _missing(wheel, "Fx"))

    _pill(slide, "Drive Tyres Comparison", 6.85, 0.44, 2.97, 0.39, _R_NAVY, 16)
    for comp, head_y, y in (("Fx", 0.47, 0.9), ("My", 3.6, 3.95)):
        _label(slide, comp, 9.9, head_y, 0.8, 0.3, size=16)
        for side, wheel, x in zip(("Left", "Right"), drive_wheels(spec.drive_a), (6.8, 10.05)):
            _result_plot(slide, figure_path(config, f"pv_drive_{comp}_{side}"), x, y,
                         PV_SLOT_IN, side, _missing(wheel, comp))

    _note_box(slide, [_OBSERVATION], 0.15, 6.86, 8.9, 0.52)
    overall = f"Overall - {spec.route}:" if spec.route else "Overall:"
    _note_box(slide, [overall, "•  …", "•  …", "•  …"], 9.25, 6.49, 3.6, 0.97)


def _cv_slide(deck: "_Deck", spec, config, rows) -> None:
    from dtt.reporting.result_figures import CV_SLOT_IN, figure_path
    slide = deck.result()
    subtitle = spec.subtitle or (f"Diff Btw {spec.label_b} vs {spec.label_a}"
                                 if spec.compared else spec.label_a)
    _runs_box(slide, [[("Force Distribution ", 20, True, _BLACK, _R_FONT_CV),
                       (subtitle, 14, True, _BLACK, _R_FONT_CV)]], 0.13, 0.02, 13.0, 0.45)
    cols = (("Fz", 1.5), ("Fx", 5.45), ("Fy", 9.3))
    for comp, x in cols:
        _label(slide, comp, x, 0.45, CV_SLOT_IN[0], 0.3, size=16, align=PP_ALIGN.CENTER)
    if len(rows) > 1:
        _rule(slide, 0.17, 3.75, 13.27, 3.75)
    for row, y in zip(rows, (0.85, 3.9)):
        names = [(f"{spec.label_a}", 14, True, _R_CYAN, _R_FONT_CV)]
        if spec.compared:
            names = [(f"{spec.label_a} vs ", 14, True, _R_CYAN, _R_FONT_CV),
                     (spec.label_b, 14, True, _R_BLUE, _R_FONT_CV)]
        _runs_box(slide, [[(f"{row.title} Force dstrb’n", 14, True, _BLACK, _R_FONT_CV)], names],
                  0.15, y + 0.4, 1.3, 1.6)
        for comp, x in cols:
            _result_plot(slide, figure_path(config, f"cv_{comp}_axle{row.axle}"), x, y,
                         CV_SLOT_IN, "", _missing(row.wheel, comp))
            # an empty, unfilled box inside the plot for the analyst's note
            _runs_box(slide, [[("", 12, False, _BLACK, _R_FONT_CV)]],
                      x + 0.6, y + 0.55, 1.5, 0.9)


def _cell_border(cell, color: str = "000000", width_emu: int = 9525) -> None:
    from pptx.oxml.ns import qn
    from lxml import etree
    tcPr = cell._tc.get_or_add_tcPr()
    # schema order: the four borders come before the cell fill
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag), w=str(width_emu))
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=color)
        tcPr.insert(i, ln)


def _tyre_life_table(slide, spec) -> None:
    """Empty "Projected tyre life" grid for the analyst: Front/Rear rows,
    Left/Right per dataset. The tool does not compute tyre life."""
    names = [spec.label_a] + ([spec.label_b] if spec.compared else [])
    ncols = 1 + 3 * len(names) - 1          # label, then Left/Right per dataset, spacer between
    col_w, row_h = 0.56, 0.25
    left = 11.93 - ncols * col_w
    table = slide.shapes.add_table(5, ncols, Inches(left), Inches(0.02),
                                   Inches(ncols * col_w), Inches(5 * row_h)).table
    table.first_row = False
    table.horz_banding = False
    for c in range(ncols):
        table.columns[c].width = Inches(col_w)
    spacers = {3 * k for k in range(1, len(names))}       # gap column between datasets
    for r in range(5):
        table.rows[r].height = Inches(row_h)
        for c in range(ncols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _WHITE
            cell.margin_left = cell.margin_right = Inches(0.02)
            cell.margin_top = cell.margin_bottom = 0
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            # an empty cell is sized by its end-of-paragraph font, 18 pt by default
            cell.text_frame.paragraphs[0]._p.get_or_add_endParaRPr().set("sz", "900")
            if not ((c == 0 and r <= 2) or c in spacers):
                _cell_border(cell)

    def put(r, c, text, bold=False, size=9):
        cell = table.cell(r, c)
        cell.text = text
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Calibri"
            run.font.color.rgb = _BLACK

    table.cell(0, 1).merge(table.cell(0, ncols - 1))
    put(0, 1, "Projected tyre life", bold=True, size=11)
    for k, name in enumerate(names):
        c = 1 + 3 * k
        table.cell(1, c).merge(table.cell(1, c + 1))
        put(1, c, name)
        put(2, c, "Left")
        put(2, c + 1, "Right")
    put(3, 0, "Front", bold=True)
    put(4, 0, "Rear", bold=True)


def _combination_slide(deck: "_Deck", spec, config) -> None:
    from dtt.reporting.result_figures import COMBO_CBAR_IN, COMBO_IN, figure_path
    slide = deck.result()
    _runs_box(slide, [[("Combination of Fx & Fy", 20, True, _BLACK, _R_FONT_CV),
                       (" – All wheels", 14, True, _BLACK, _R_FONT_CV)]],
              0.17, 0.04, 7.8, 0.4)
    vs = f"{spec.label_a} Vs {spec.label_b}" if spec.compared else spec.label_a
    _runs_box(slide, [[(f"{vs} - Heatmap of Fx & Fy", 16, False, _R_MAGENTA, _R_FONT_CV)]],
              0.2, 0.42, 8.2, 0.37)

    w, h = COMBO_IN
    if spec.compared:
        halves = (("combo_a", spec.label_a, 0.35), ("combo_b", spec.label_b, 6.7))
        cbar_x = 12.58
        _rule(slide, 6.67, 0.82, 6.67, 6.73, color=_R_DIVIDE, width_pt=3.0)
    else:
        halves = (("combo_a", spec.label_a, 3.4),)
        cbar_x = 3.4 + w + 0.05
    for name, label, x in halves:
        _insert_image_or_message(slide, figure_path(config, name), Inches(x), Inches(1.55),
                                 Inches(w), Inches(h),
                                 reason="Not available: no Fx/Fy pair was recorded.")
        _label(slide, label, x, 6.36, w, 0.35, size=18, align=PP_ALIGN.CENTER)
    cbar = figure_path(config, "combo_cbar")
    if cbar.exists():
        _add_picture_fit(slide, cbar, Inches(cbar_x), Inches(1.75),
                         Inches(COMBO_CBAR_IN[0]), Inches(COMBO_CBAR_IN[1]))
    _tyre_life_table(slide, spec)              # last, so it sits above the pictures

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.64), Inches(6.76),
                                 Inches(11.69), Inches(0.67))
    box.fill.solid()
    box.fill.fore_color.rgb = _R_BOX
    box.line.color.rgb = _R_MAGENTA
    box.line.width = Pt(2.25)
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    box.line.dash_style = MSO_LINE_DASH_STYLE.SQUARE_DOT
    tf = box.text_frame
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i in range(2):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = f"•   {_OBSERVATION}"
        r.font.size, r.font.name = Pt(14), _R_FONT
        r.font.color.rgb = _BLACK


def _result_slides(deck: "_Deck", spec, config) -> None:
    """The format's section slide, then its force-distribution pages."""
    if spec.fmt == "pv":
        deck.section("PV Format")
        _pv_page_fy_fz(deck, spec, config)
        _pv_page_fx_my(deck, spec, config)
    else:
        deck.section("CV Format")
        rows = spec.cv_rows or []
        if not rows:
            slide = deck.result()
            _label(slide, "Force Distribution", 0.13, 0.02, 12.0, 0.45, size=20, font=_R_FONT_CV)
            _add_textbox(slide, "Not available: no wheel positions were recognised in this "
                                "study.", L, T, IMG_W, Inches(0.5), font_size=12, italic=True,
                         color=_GREY)
        for i in range(0, len(rows), 2):
            _cv_slide(deck, spec, config, rows[i:i + 2])


def _boxplot_slides(deck: "_Deck", figs: Path) -> None:
    """Box plots unchanged, under the result slides' title style and footer."""
    for group, title in ((("Fx", "Fy", "Fz"), "Force Distribution"),
                         (("Mx", "My", "Mz"), "Moment Distribution")):
        present = [c for c in group if (figs / f"boxplot_{c}.png").exists()]
        if not present and group[0] == "Mx":
            continue                      # a forces-only study has no moment slide
        slide = deck.result()
        _label(slide, f"Box Plots – {title} ({' / '.join(group)})", 0.13, 0.02, 12.5, 0.5,
               size=28, color=_R_TITLE)
        _label(slide, "Median, quartiles, P5 / mean / P95 per wheel", 0.15, 0.62, 12.0, 0.3,
               size=14, color=_GREY)
        third_w = int(IMG_W) // 3
        for i, comp in enumerate(group):
            _insert_image_or_message(slide, figs / f"boxplot_{comp}.png",
                                     L + i * third_w, Inches(1.05), third_w, Inches(5.6))
    if (figs / "boxplot_parameters.png").exists():
        slide = deck.result()
        _label(slide, "Box Plots – Vehicle & Other Parameters", 0.13, 0.02, 12.5, 0.5,
               size=28, color=_R_TITLE)
        _label(slide, "Every other measured channel, each on its own unit", 0.15, 0.62,
               12.0, 0.3, size=14, color=_GREY)
        _insert_image_or_message(slide, figs / "boxplot_parameters.png", L, Inches(1.05),
                                 IMG_W, Inches(5.6))


# ------------------------------------------------------------------- content

def _processing_lines(config: RunConfig, metadata: Dict) -> List[str]:
    """What was done to the data, in the order it was done -- so the report
    says how its numbers were produced, not only what they are."""
    lines = []
    if getattr(config, "famos_mode", True):
        lines.append("FAMOS recipe: forces/moments smo(0.1 s); Latacc FiltLP(4, 5 Hz) "
                     "→ smo(0.5 s); Longacc, speed smo(0.5 s); red(10)")
    else:
        lines.append(f"Butterworth low-pass {getattr(config, 'filter_cutoff', '?')} Hz, "
                     f"order {getattr(config, 'filter_order', '?')}")
    lines.append("Artefact blanking: samples beyond the P1–P99 band ± one band-width")
    stop = metadata.get("stop_removal") or {}
    if getattr(config, "remove_stops", False):
        lines.append(f"Stops removed: {stop.get('stops_removed', 0)} "
                     f"({stop.get('stop_seconds', 0):.0f} s) from every analysis")
    else:
        lines.append("Stops kept (standstill is in every statistic)")
    if getattr(config, "despike", False):
        lines.append("Physical despike: on")
    return lines


def _processing_summary(config: RunConfig, metadata: Dict) -> str:
    return "Processing:\n" + "\n".join(f"  {l}" for l in _processing_lines(config, metadata))


def _add_sessions_slide(deck: _Deck, sessions: Dict, figs: Path) -> None:
    from dtt.analysis.sessions_plot import SESSIONS_FIGURE, session_spans
    slide = deck.content(f"Recording Sessions  –  {sessions['sessions']} joined",
                         "Gaps between sessions removed · elapsed time continuous · "
                         "each session conditioned on its own before joining")
    rows = [f"Session {i}:  {Path(src).name or src or '—'}   "
            f"{s:.0f} – {e:.0f} s  ({(e - s) / 60:.1f} min)"
            for i, ((s, e), src) in enumerate(
                zip(session_spans(sessions), sessions["session_sources"]), start=1)]
    text_h = Emu(240000 * (len(rows) + 1))
    _add_textbox(slide, "\n".join(rows), L, T, IMG_W, text_h, font_size=13)
    _insert_image_or_message(slide, figs / SESSIONS_FIGURE, L, T + text_h,
                             IMG_W, IMG_H - text_h)


def build_report(
    config: RunConfig,
    validation_report: ValidationReport,
    stats: Dict,
    metadata: Dict,
    analysis_only: bool = False,
    result_spec=None,
) -> Path:
    """``result_spec`` is what :func:`dtt.reporting.result_figures.generate_result_figures`
    returned for this run -- it knows whether the comparison study could
    actually be read. Omitted, it is resolved from ``config``.

    Slide order: cover and data summary; the result slides in the charter
    format (PV or CV section, its force-distribution pages); box plots; the
    Fx & Fy combination heatmap; then the remaining analyses.

    ``analysis_only=True`` marks the report as produced by
    ``workflow_mode="analysis"`` — no fresh preprocessing/validation pass, the
    channel-validation and file/duration/units figures describe the
    *existing* processed_data.csv this run read, not a run that just
    ingested and conditioned it. Made visible on the cover and overview so a
    reduced-metadata analysis-only report can't be mistaken for a full run.
    """
    figs = config.figures_dir

    # This study's own axle-generalized channel model, not the legacy fixed
    # FL/FR/RL/RR constants -- a recording with only FR/RR (or a 6-wheel
    # truck) must get exactly the slides its own data supports.
    rc = config.run_channels
    wheel_groups = rc.wheel_groups if rc is not None else WHEEL_GROUPS
    stat_channels = rc.mandatory_channels if rc is not None else MANDATORY_CHANNELS
    wheels = list(wheel_groups.keys())

    now = datetime.datetime.now()
    now_str = now.strftime("%Y-%m-%d %H:%M")
    sessions = session_info(metadata)
    deck = _Deck(_new_presentation(),
                 footer=f"WFT Road Load Report  ·  {config.vehicle_name}  ·  "
                        f"{config.study_name}")

    # ---- cover
    cover_lines = [f"{config.vehicle_name}  ·  Study {config.study_name}",
                   now.strftime("%d %B %Y")]
    if analysis_only:
        cover_lines.append("ANALYSIS-ONLY MODE")
    deck.cover("WFT Road Load Data – Tyre Duty Cycle Report", cover_lines)

    # ---- data & processing
    deck.divider("Data & Processing", "Recording · validation · statistics")

    slide = deck.content("Study Overview", f"{config.vehicle_name}  ·  {now_str}")
    overview = [
        ("Study", config.study_name),
        ("Source", str(metadata.get("file_name", "N/A"))),
        ("Duration", f"{metadata.get('duration_s', 0):,.1f} s  "
                     f"({metadata.get('duration_s', 0) / 60:.1f} min)"),
        ("Records", f"{metadata.get('rows', 0):,} rows × {metadata.get('columns', 0)} channels"),
        ("Sampling rate", f"{metadata.get('sampling_rate_hz', 0):.0f} Hz"),
        ("Units", "N → daN (÷10)" if metadata.get("n_to_dan_applied") else "daN (unchanged)"),
        ("Wheel positions", ", ".join(wheels) or "—"),
        ("Processing", "\n".join(_processing_lines(config, metadata))),
    ]
    if sessions:
        overview.append(("Sessions",
                         f"{sessions['sessions']} recordings joined end to end — seams at "
                         f"{', '.join(f'{t:.0f}' for t in sessions['seam_times_s'])} s"))
    if analysis_only:
        overview.append(("Mode", "ANALYSIS-ONLY: generated from an existing "
                                 "processed_data.csv, no fresh preprocessing or "
                                 "validation pass; see the pipeline log for "
                                 "provenance warnings"))
    _add_key_value_table(slide, overview, L, T, IMG_W, IMG_H)

    if sessions:
        _add_sessions_slide(deck, sessions, figs)

    slide = deck.content(
        "Channel Validation Summary",
        f"Missing: {len(validation_report.missing_channels)}   ·   "
        f"Optional present: {len(validation_report.optional_present)}   ·   "
        f"Total NaN rows: {validation_report.total_nan_rows:,}")
    _add_validation_table(slide, validation_report, L, T, IMG_W, IMG_H)

    slide = deck.content("Statistical Summary  –  P80 / P90 / P95",
                         "Per wheel force channel, daN")
    _add_stats_table(slide, stats, stat_channels, L, T, IMG_W, IMG_H)

    # ---- result slides (charter format), box plots, combination heatmap
    if result_spec is None:
        from dtt.reporting.result_figures import resolve_spec
        result_spec = resolve_spec(config)
    _result_slides(deck, result_spec, config)
    _boxplot_slides(deck, figs)
    _combination_slide(deck, result_spec, config)

    # ---- module 1
    deck.divider("Module 1", "Tyre load distributions & heatmaps")

    half_h = (int(IMG_H) - int(Inches(0.1))) // 2
    for wheel in wheels:
        slide = deck.content(f"Histograms  –  {wheel}  (Distance & Percentage)",
                             "Share of distance travelled, then share of samples, per load bin")
        _insert_image_or_message(slide, figs / f"hist_distance_{wheel}.png", L, T, IMG_W, half_h)
        _insert_image_or_message(slide, figs / f"hist_percentage_{wheel}.png",
                                 L, T + half_h + Inches(0.1), IMG_W, half_h)

    for wheel in wheels:
        slide = deck.content(f"AUC Distribution  –  {wheel}  (Distance & Percentage)",
                             "Load density curves with P5 / P95")
        _insert_image_or_message(slide, figs / f"auc_distance_{wheel}.png", L, T, IMG_W, half_h)
        _insert_image_or_message(slide, figs / f"auc_percentage_{wheel}.png",
                                 L, T + half_h + Inches(0.1), IMG_W, half_h)

    for hm_name, hm_title in [
        ("heatmap_fx_fy_all",  "Heatmap  –  Fx vs Fy  (All Wheels)"),
        ("heatmap_fz_fy_all",  "Heatmap  –  Fz vs Fy  (All Wheels)"),
        ("heatmap_fz_fx_all",  "Heatmap  –  Fz vs Fx  (All Wheels)"),
    ]:
        slide = deck.content(hm_title, "Force combination occurrence")
        _insert_image_or_message(slide, figs / f"{hm_name}.png", L, T, IMG_W, IMG_H)

    slide = deck.content("Hexbin Density  –  Fy vs Fx  (All Wheels)",
                         "Sample density, log scale")
    for (lf, tf, gw, gh), wheel in zip(_grid_positions(len(wheels), L, T, IMG_W, IMG_H), wheels):
        _insert_image_or_message(slide, figs / f"heatmap_hexbin_{wheel}.png", lf, tf, gw, gh)

    # ---- conclusions
    slide = deck.content("Engineering Conclusions", f"{config.vehicle_name}  ·  {now_str}")
    present_ch = [ch for ch, st in validation_report.channel_status.items() if st == "PRESENT"]
    missing_ch = validation_report.missing_channels
    results = [f"Channels validated: {len(present_ch)} / {len(stat_channels)} mandatory",
               f"Total records: {metadata.get('rows', 0):,}",
               f"Study duration: {metadata.get('duration_s', 0):,.1f} s"]
    if stats:
        fz_keys = [ch for ch in stats if "Fz" in ch and stats[ch].get("P95") is not None]
        if fz_keys:
            top = max(fz_keys, key=lambda ch: stats[ch]["P95"])
            results.append(f"Max Fz P95 (all wheels): {stats[top]['P95']:.1f} daN "
                           f"({display_name(top)})")
    notes = []
    if missing_ch:
        notes.append(f"Missing channels: {', '.join(missing_ch)}")
    if sessions:
        notes.append(f"Joined record: {sessions['sessions']} sessions, seams at "
                     f"{', '.join(f'{t:.0f}' for t in sessions['seam_times_s'])} s")
    if analysis_only:
        notes.append("ANALYSIS-ONLY MODE — no fresh preprocessing/validation pass; "
                     "generated from an existing processed_data.csv")
    notes.extend(_processing_lines(config, metadata))
    notes.append("Refer to individual slides for channel-level analysis")
    half_w = (int(IMG_W) - int(Inches(0.3))) // 2
    _add_panel(slide, "Key Results", results, L, T, half_w, IMG_H)
    _add_panel(slide, "Notes & Processing", notes, L + half_w + Inches(0.3), T,
               half_w, IMG_H)

    prs = deck.prs
    date_str = now.strftime("%Y%m%d")
    out_fname = (config.run_output_dir / f"WFT_Report_{config.vehicle_name}_{date_str}.pptx").resolve()
    prs.save(str(out_fname))

    blank_slides = _audit_slides(prs)
    if blank_slides:
        logger.warning("Report has %d title-only slide(s) with no content or "
                       "explanation: %s", len(blank_slides), "; ".join(blank_slides))
    logger.info("PowerPoint report saved: %s  (%d slides, %d wheel(s): %s, "
                "stats for %d/%d channels)",
                out_fname, len(prs.slides), len(wheels), ", ".join(wheels),
                len([ch for ch in stat_channels if ch in stats]), len(stat_channels))
    return out_fname
