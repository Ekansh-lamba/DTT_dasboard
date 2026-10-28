"""The charter-format result slides (dtt/reporting/result_figures.py and the
result-slide builders in report_builder.py).

What is pinned here:
- the slides come in the reference deck's order (format section, the
  force-distribution pages, box plots, the Fx & Fy combination heatmap) ahead
  of the remaining analyses;
- every title, heading, wheel label, pill and summary box is a native,
  editable text shape -- the only pictures are the plots;
- a CV study gets one row per recorded axle, titled from the vehicle preset;
- a single-study run still gets the same layouts, drawn for dataset A alone;
- the stats strip's numbers are the Statistics stage's numbers.

Small synthetic frames, built here -- no recording, no full pipeline run.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from pptx import Presentation

from dtt.analysis.plot_style import stats_strip_lines, stats_strip_rows
from dtt.analysis.statistics import channel_statistics
from dtt.config import RunConfig
from dtt.reporting.report_builder import _LAYOUT_RESULT, _audit_slides, build_report
from dtt.reporting.result_figures import (figure_path, generate_result_figures,
                                          resolve_spec)
from dtt.run_channels import build_run_channels
from dtt.validation.validator import ValidationReport

pytestmark = pytest.mark.filterwarnings("ignore")

_N = 3000


def _car_frame(seed: int, fz_rear: float = 700.0) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    out = {"Time": np.arange(_N) * 0.01, "Vehicle_Speed": np.full(_N, 40.0)}
    for w in ("FL", "FR", "RL", "RR"):
        out[f"{w}_Fx"] = rng.normal(10, 40, _N)
        out[f"{w}_Fy"] = rng.normal(0, 30, _N)
        out[f"{w}_Fz"] = rng.normal(fz_rear if w[0] == "R" else 650.0, 60, _N)
        out[f"{w}_My"] = rng.normal(0, 150, _N)
    return pd.DataFrame(out)


def _truck_frame(seed: int) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    out = {"Time": np.arange(_N) * 0.01}
    for axle, fz in ((1, 6000.0), (2, 8000.0), (3, 7000.0)):
        for side in "LR":
            out[f"A{axle}{side}_Fx"] = rng.normal(50, 200, _N)
            out[f"A{axle}{side}_Fy"] = rng.normal(0, 150, _N)
            out[f"A{axle}{side}_Fz"] = rng.normal(fz, 600, _N)
    return pd.DataFrame(out)


def _config(tmp_path: Path, df: pd.DataFrame, name: str, **kw) -> RunConfig:
    cfg = RunConfig(output_dir=tmp_path, study_name=name, vehicle_name=name, **kw)
    cfg.run_channels = build_run_channels(list(df.columns))
    return cfg


def _finished_study(tmp_path: Path, df: pd.DataFrame, name: str) -> Path:
    d = tmp_path / name
    d.mkdir(parents=True, exist_ok=True)
    df.to_csv(d / "processed_data.csv", index=False)
    return d


def _report(cfg: RunConfig, df: pd.DataFrame) -> Path:
    spec = generate_result_figures(df, cfg)
    stats = channel_statistics(df)
    vr = ValidationReport(file_name=cfg.study_name, rows=len(df), columns=len(df.columns),
                          duration_s=_N / 100, sampling_rate_hz=100.0,
                          channel_status={c: "PRESENT" for c in df.columns if c != "Time"})
    meta = {"file_name": cfg.study_name, "rows": len(df), "columns": len(df.columns),
            "duration_s": _N / 100, "sampling_rate_hz": 100.0}
    return build_report(cfg, vr, stats, meta, result_spec=spec)


def _titles(prs) -> list:
    out = []
    for slide in prs.slides:
        first = next((sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text),
                     None)
        out.append(first.text_frame.text.split("\n")[0] if first else "")
    return out


def _texts(slide) -> list:
    return [sh.text_frame.text.strip() for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]


# ------------------------------------------------------------- PV, compared

@pytest.fixture(scope="module")
def pv_report(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pv")
    df_a = _car_frame(1, fz_rear=760.0)
    study_b = _finished_study(tmp, _car_frame(2, fz_rear=560.0), "ic_study")
    cfg = _config(tmp, df_a, "ev_study", report_compare_study=study_b,
                  report_label_a="EV", report_label_b="IC", report_route="City",
                  report_drive_axle_a="rear", report_drive_axle_b="front")
    return cfg, Presentation(str(_report(cfg, df_a)))


def test_pv_slides_come_in_the_reference_order(pv_report):
    _, prs = pv_report
    titles = _titles(prs)
    start = titles.index("PV Format")
    assert titles[start - 1].startswith("Statistical Summary")
    assert titles[start + 1] == "Force Distribution - EV vs IC – Fy & Fz"
    assert titles[start + 2] == "Force Distribution - EV vs IC – Fx"
    assert titles[start + 3].startswith("Box Plots – Force Distribution")
    combo = titles.index("Combination of Fx & Fy – All wheels")
    assert combo > start + 3
    assert all(t.startswith("Box Plots") for t in titles[start + 3:combo])
    assert titles[combo + 1] == "Module 1"


def test_pv_text_is_native_and_only_plots_are_pictures(pv_report):
    cfg, prs = pv_report
    titles = _titles(prs)
    page1 = prs.slides[titles.index("PV Format") + 1]
    page2 = prs.slides[titles.index("PV Format") + 2]
    assert page1.slide_layout.name == _LAYOUT_RESULT

    t1 = _texts(page1)
    for expected in ("City", "Fy", "Fz", "FL", "FR", "RL", "RR", "Key observation: …"):
        assert expected in t1, (expected, t1)
    assert sum(1 for sh in page1.shapes if sh.shape_type == 13) == 8      # 2 x 2 x (Fy, Fz)

    t2 = _texts(page2)
    for expected in ("By Wheel Position", "Drive Tyres Comparison", "Fx", "My",
                     "Left", "Right", "Key observation: …"):
        assert expected in t2, (expected, t2)
    assert any(t.startswith("Overall - City:") for t in t2)
    assert sum(1 for sh in page2.shapes if sh.shape_type == 13) == 8      # 4 Fx + 2 Fx + 2 My

    # every picture on the result slides is one of the generated plot files
    report_dir = Path(cfg.figures_dir) / "report"
    assert len(list(report_dir.glob("pv_*.png"))) == 16


def test_combination_slide_has_editable_table_and_summary_box(pv_report):
    _, prs = pv_report
    titles = _titles(prs)
    slide = prs.slides[titles.index("Combination of Fx & Fy – All wheels")]
    texts = _texts(slide)
    assert "EV Vs IC - Heatmap of Fx & Fy" in texts
    assert "EV" in texts and "IC" in texts
    assert any("Key observation" in t for t in texts)
    tables = [sh.table for sh in slide.shapes if sh.has_table]
    assert len(tables) == 1
    cells = [c.text for row in tables[0].rows for c in row.cells]
    for expected in ("Projected tyre life", "EV", "IC", "Left", "Right", "Front", "Rear"):
        assert expected in cells
    pics = sum(1 for sh in slide.shapes if sh.shape_type == 13)
    assert pics == 3                     # EV heatmaps, IC heatmaps, shared colorbar


def test_pv_drive_comparison_pairs_each_datasets_driven_axle(pv_report):
    """EV rear-driven against IC front-driven: Left is RL (EV) vs FL (IC)."""
    cfg, _ = pv_report
    spec = resolve_spec(cfg)
    assert (spec.drive_a, spec.drive_b) == ("rear", "front")
    assert figure_path(cfg, "pv_drive_Fx_Left").exists()
    assert figure_path(cfg, "pv_drive_My_Right").exists()


def test_no_title_only_slides(pv_report):
    _, prs = pv_report
    assert _audit_slides(prs) == []


# ------------------------------------------------------------------- CV

def test_cv_gets_one_row_per_axle_titled_from_the_preset(tmp_path):
    df_a = _truck_frame(3)
    study_b = _finished_study(tmp_path, _truck_frame(4), "mrf_study")
    cfg = _config(tmp_path, df_a, "apollo_study", report_compare_study=study_b,
                  report_label_a="Apollo", report_label_b="MRF")
    spec = resolve_spec(cfg)
    assert spec.fmt == "cv"
    assert [(r.axle, r.title) for r in spec.cv_rows] == [
        (1, "Steer Axle"), (2, "Drive Axle"), (3, "Tag Axle")]

    prs = Presentation(str(_report(cfg, df_a)))
    titles = _titles(prs)
    start = titles.index("CV Format")
    cv_slides = [prs.slides[start + 1], prs.slides[start + 2]]        # 3 axles -> 2 + 1 rows
    texts = _texts(cv_slides[0])
    assert texts[0].startswith("Force Distribution Diff Btw MRF vs Apollo")
    for expected in ("Fz", "Fx", "Fy"):
        assert expected in texts
    assert any(t.startswith("Steer Axle Force dstrb") and "Apollo vs MRF" in t for t in texts)
    assert any(t.startswith("Drive Axle Force dstrb") for t in texts)
    assert any(t.startswith("Tag Axle Force dstrb") for t in _texts(cv_slides[1]))
    assert sum(1 for sh in cv_slides[0].shapes if sh.shape_type == 13) == 6
    # an empty note box sits in every plot for the analyst
    empty = [sh for sh in cv_slides[0].shapes
             if sh.has_text_frame and not sh.text_frame.text.strip() and sh.shape_type == 17]
    assert len(empty) == 6


# ---------------------------------------------------------- single study

def test_single_study_draws_the_same_layouts_for_dataset_a(tmp_path):
    df = _car_frame(5)
    cfg = _config(tmp_path, df, "solo", report_label_a="EV")
    prs = Presentation(str(_report(cfg, df)))
    titles = _titles(prs)
    assert "Force Distribution - EV – Fy & Fz" in titles
    assert "Force Distribution - EV – Fx" in titles
    combo = prs.slides[titles.index("Combination of Fx & Fy – All wheels")]
    assert "EV - Heatmap of Fx & Fy" in _texts(combo)
    assert not figure_path(cfg, "combo_b").exists()


def test_unreadable_compare_study_falls_back_to_single(tmp_path):
    df = _car_frame(6)
    cfg = _config(tmp_path, df, "lonely", report_label_a="EV",
                  report_compare_study=tmp_path / "does_not_exist")
    spec = generate_result_figures(df, cfg)
    assert spec.label_b is None
    assert figure_path(cfg, "pv_Fz_FL").exists()


def test_fr_rr_study_marks_missing_wheels_instead_of_dropping_slots(tmp_path):
    rng = np.random.default_rng(7)
    cols = {"Time": np.arange(_N) * 0.01}
    for w, n in (("FR", 2), ("RR", 1)):
        for c in ("Fx", "Fy", "Fz"):
            cols[f"{w}_{c}_{n}"] = rng.normal(500 if c == "Fz" else 0, 50, _N)
    df = pd.DataFrame(cols)
    cfg = _config(tmp_path, df, "fr_rr")
    prs = Presentation(str(_report(cfg, df)))
    titles = _titles(prs)
    page1 = prs.slides[titles.index("PV Format") + 1]
    texts = _texts(page1)
    assert "FR" in texts and "RR" in texts
    assert any(t.startswith("FL Fy: not recorded") for t in texts)
    assert sum(1 for sh in page1.shapes if sh.shape_type == 13) == 4


# ------------------------------------------------------------ stats strip

def test_stats_strip_equals_the_statistics_stage():
    """Rounded to whole numbers for display, the strip's Min/Median/Mean/Max
    are the Statistics stage's own values for the same channel."""
    df = _car_frame(8)
    stats = channel_statistics(df)
    for col in ("FL_Fx", "RR_Fz"):
        _, rows = stats_strip_rows([("A", df[col].to_numpy())], decimals=0)
        row = dict(zip(["Run", "Min", "P5", "Q1", "Median", "Mean", "Q3", "P95", "Max"], rows[0]))
        s = stats[col]
        for key, stat in (("Min", "min"), ("Median", "median"), ("Mean", "mean"), ("Max", "max")):
            assert row[key] == f"{s[stat]:.0f}", (col, key)
        assert row["P95"] == f"{s['P95']:.0f}"


def test_stats_strip_lines_align_columns():
    header, rows = stats_strip_rows([("EV", np.array([1.0, 2.0, 300.0])),
                                     ("IC", np.array([-1000.0, 5.0, 9.0]))], decimals=0)
    a, b = stats_strip_lines(header, rows)
    assert a.startswith("EV | Min:") and b.startswith("IC | Min:")
    assert [i for i, ch in enumerate(a) if ch == "|"] == [i for i, ch in enumerate(b) if ch == "|"]


# ------------------------------------------------------------ GUI -> CLI

def test_gui_request_passes_report_options_to_the_cli():
    """The GUI shells out to the CLI, so a report field only works if its
    flag reaches the command line -- and a blank field must add nothing."""
    pytest.importorskip("PySide6")
    from gui.workers.pipeline_worker import RunRequest

    plain = RunRequest(csv_path=Path("x.csv")).to_cmd()
    assert not any(a.startswith(("--report", "--compare", "--label", "--route", "--drive"))
                   for a in plain)

    cmd = RunRequest(csv_path=Path("x.csv"), report_format="pv",
                     compare_study=Path("outputs/ic"), label_a="EV", label_b="IC",
                     route="City", drive_axle_a="rear", drive_axle_b="front").to_cmd()
    pairs = dict(zip(cmd, cmd[1:]))
    assert pairs["--report-format"] == "pv"
    assert Path(pairs["--compare-study"]) == Path("outputs/ic")
    assert (pairs["--label-a"], pairs["--label-b"], pairs["--route"]) == ("EV", "IC", "City")
    assert (pairs["--drive-axle-a"], pairs["--drive-axle-b"]) == ("rear", "front")


def test_stats_strip_rows_default_rounding_unchanged():
    _, rows = stats_strip_rows([("A", np.array([1.25, 2.5, 3.75]))])
    assert rows[0][1] == "1.2" or rows[0][1] == "1.3"       # one decimal, as before
    assert "." in rows[0][4]
