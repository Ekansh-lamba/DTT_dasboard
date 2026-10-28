"""Report-generation coverage for dtt/reporting/report_builder.py.

Before this file, report_builder.py had zero test coverage anywhere in the
repo -- which is exactly how it drifted onto the legacy hardcoded
WHEEL_GROUPS/MANDATORY_CHANNELS constants while every other analysis module
(statistics.py, histograms.py, heatmaps.py, boxplots.py) had
already moved to the dynamic config.run_channels model. A real FR/RR-only
recording (no FL/RL) produced: an empty statistics slide (legacy channel
names never matched the study's real, axle-suffixed stats keys) and
title-only blank FL/RL slides (legacy wheel list always assumed all four).

These tests build report_builder's own inputs by hand -- no full pipeline
run, no real recording -- fast and deterministic, matching the rest of this
suite. The one real-data end-to-end check lives outside pytest (see the
session notes), since running the full pipeline is much too slow to
run on every `pytest tests/`.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from dtt.config import RunConfig
from dtt.reporting.report_builder import (
    _add_stats_table,
    _audit_slides,
    _grid_positions,
    _insert_image_or_message,
    build_report,
)
from dtt.run_channels import build_run_channels
from dtt.validation.validator import ValidationReport

pytestmark = pytest.mark.filterwarnings("ignore")


def _fr_rr_config(tmp_path: Path, study_name: str = "fr_rr_study") -> RunConfig:
    """A RunConfig shaped like the real 'dtt data/' recordings: only FR/RR
    wheel positions, output isolated under tmp_path (auto-created by
    RunConfig.__post_init__, torn down by pytest -- nothing touches the
    repo or dtt/outputs/)."""
    cfg = RunConfig(output_dir=tmp_path, study_name=study_name)
    cfg.run_channels = build_run_channels(
        ["FR_Fx_2", "FR_Fy_2", "FR_Fz_2", "RR_Fx_1", "RR_Fy_1", "RR_Fz_1"])
    return cfg


def _fr_rr_stats() -> dict:
    entry = lambda: {"count": 1000, "mean": 10.0, "median": 9.5, "std": 2.0,
                     "min": 1.0, "max": 30.0, "P80": 15.0, "P90": 18.0, "P95": 22.0}
    return {"FR_Fx_2": entry(), "FR_Fy_2": entry(), "FR_Fz_2": entry(),
           "RR_Fx_1": entry(), "RR_Fy_1": entry(), "RR_Fz_1": entry()}


def _fr_rr_validation_report() -> ValidationReport:
    return ValidationReport(
        file_name="fr_rr_study", rows=1000, columns=8, duration_s=10.0,
        sampling_rate_hz=100.0,
        channel_status={ch: "PRESENT" for ch in
                        ("FR_Fx_2", "FR_Fy_2", "FR_Fz_2", "RR_Fx_1", "RR_Fy_1", "RR_Fz_1")},
    )


def _fr_rr_metadata() -> dict:
    return {"file_name": "fr_rr_study", "rows": 1000, "columns": 8,
           "duration_s": 10.0, "sampling_rate_hz": 100.0, "n_to_dan_applied": True}


def _slide_titles(pptx_path: Path) -> list[str]:
    prs = Presentation(str(pptx_path))
    titles = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text:
                titles.append(sh.text_frame.text.split("\n")[0])
                break
    return titles


# --------------------------------------------- missing wheel positions

def test_report_only_generates_slides_for_wheels_present(tmp_path):
    """FR/RR-only data must never produce an FL or RL slide -- the legacy
    WHEEL_GROUPS constant always assumed all four."""
    cfg = _fr_rr_config(tmp_path)
    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())

    titles = _slide_titles(out)
    assert not any("FL" in t for t in titles), titles
    assert not any(" RL " in t or t.endswith("RL") for t in titles), titles
    assert any("FR" in t for t in titles)
    assert any("RR" in t for t in titles)


def test_report_slide_count_matches_wheel_count_not_legacy_four(tmp_path):
    cfg = _fr_rr_config(tmp_path)
    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())

    titles = _slide_titles(out)
    histogram_slides = [t for t in titles if t.startswith("Histograms")]
    auc_slides       = [t for t in titles if t.startswith("AUC Distribution")]
    assert len(histogram_slides) == 2   # FR, RR -- not the legacy 4
    assert len(auc_slides) == 2


# ------------------------------------- severity & fatigue removed (2026-10-28)

def _png_1x1() -> bytes:
    import base64
    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42Y"
        "AAAAASUVORK5CYII=")


def test_no_severity_slide_even_when_its_table_exists(tmp_path):
    """G-severity and DLC are still computed and saved (severity_summary.json,
    severity_table.png) but are no longer a report slide."""
    cfg = _fr_rr_config(tmp_path)
    (cfg.figures_dir / "severity_table.png").write_bytes(_png_1x1())

    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())
    titles = _slide_titles(out)
    assert not any("Severity" in t or "G-severity" in t for t in titles), titles


def test_no_fatigue_or_psd_section_even_when_old_figures_exist(tmp_path):
    """An old study's rainflow/PSD figures are ignored: no divider, no slide."""
    cfg = _fr_rr_config(tmp_path)
    for wheel in ("FR", "RR"):
        (cfg.figures_dir / f"rainflow_{wheel}.png").write_bytes(_png_1x1())
        (cfg.figures_dir / f"psd_{wheel}.png").write_bytes(_png_1x1())

    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())
    text = " ".join(_slide_titles(out)).lower()
    for word in ("rainflow", "psd", "fatigue", "welch", "miner"):
        assert word not in text, word
    assert _slide_titles(out)[-1] == "Engineering Conclusions"


def test_slide_numbers_run_without_gaps(tmp_path):
    """Every numbered slide carries its own position -- no hole is left
    where the removed slides used to be."""
    cfg = _fr_rr_config(tmp_path)
    prs = Presentation(str(build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(),
                                        _fr_rr_metadata())))
    for i, slide in enumerate(prs.slides, start=1):
        numbers = [sh.text_frame.text for sh in slide.shapes
                   if sh.has_text_frame and sh.text_frame.text.strip().isdigit()]
        if numbers:
            assert numbers[-1] == str(i), (i, numbers)


# --------------------------------------------------- statistical summary

def test_stats_table_uses_this_studys_channel_names_not_legacy_constant(tmp_path):
    cfg = _fr_rr_config(tmp_path)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.util import Emu
    _add_stats_table(slide, _fr_rr_stats(), cfg.run_channels.mandatory_channels,
                     Emu(0), Emu(0), Emu(9000000), Emu(5000000))
    tables = [sh for sh in slide.shapes if sh.has_table]
    assert len(tables) == 1
    table = tables[0].table
    assert len(table.rows) == 1 + 6   # header + 6 real channels
    # rows are this study's channels (FR_Fx_2 ...), shown under the one
    # naming convention: the standard name, or the user's display name
    from dtt.channel_names import display_name
    assert table.cell(1, 0).text == display_name("FR_Fx_2")
    assert display_name("FR_Fx_2", names={}) == "FR_Fx"


def test_stats_table_handles_missing_fields_safely(tmp_path):
    """A channel missing a field (e.g. 'std' never computed) must show N/A,
    not raise a KeyError and not fabricate a number."""
    stats = _fr_rr_stats()
    del stats["FR_Fx_2"]["std"]
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.util import Emu
    _add_stats_table(slide, stats, list(stats.keys()),
                     Emu(0), Emu(0), Emu(9000000), Emu(5000000))
    table = [sh for sh in slide.shapes if sh.has_table][0].table
    header = [table.cell(0, c).text for c in range(len(table.columns))]
    std_col = header.index("Std")
    assert table.cell(1, std_col).text == "N/A"


def test_stats_table_empty_dataset_shows_message_not_blank(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.util import Emu
    _add_stats_table(slide, {}, ["FR_Fx_2"], Emu(0), Emu(0), Emu(9000000), Emu(5000000))
    tables = [sh for sh in slide.shapes if sh.has_table]
    texts  = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert not tables
    assert any("Not available" in t for t in texts)


def test_full_report_stats_slide_is_populated_for_fr_rr_data(tmp_path):
    """Reproduces the originally-reported bug: this exact stats shape
    (generalized axle-suffixed keys) must land in the table, not vanish."""
    cfg = _fr_rr_config(tmp_path)
    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())
    prs = Presentation(str(out))
    titles = _slide_titles(out)
    stats_slide = prs.slides[next(i for i, t in enumerate(titles)
                                  if t.startswith("Statistical Summary"))]
    tables = [sh for sh in stats_slide.shapes if sh.has_table]
    assert len(tables) == 1
    assert len(tables[0].table.rows) == 1 + 6


# ------------------------------------------------------- figure handling

def test_missing_figure_file_gets_explanatory_message(tmp_path):
    from pptx.util import Emu
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    inserted = _insert_image_or_message(
        slide, tmp_path / "does_not_exist.png", Emu(0), Emu(0), Emu(1000000), Emu(1000000))
    assert inserted is False
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert any("Not available" in t for t in texts)
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    assert not pics


def test_invalid_image_path_falls_back_to_message_not_a_crash(tmp_path):
    """A path that exists but is not a valid image (corrupt/wrong content)
    must not raise out of build_report -- it must degrade to a message."""
    bad = tmp_path / "corrupt.png"
    bad.write_text("this is not a PNG file")
    from pptx.util import Emu
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    inserted = _insert_image_or_message(
        slide, bad, Emu(0), Emu(0), Emu(1000000), Emu(1000000))
    assert inserted is False
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert any("Not available" in t for t in texts)


def test_empty_zero_byte_image_treated_as_missing(tmp_path):
    empty = tmp_path / "empty.png"
    empty.write_bytes(b"")
    from pptx.util import Emu
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    inserted = _insert_image_or_message(
        slide, empty, Emu(0), Emu(0), Emu(1000000), Emu(1000000))
    assert inserted is False


# --------------------------------------------------------- layout / grid

def test_grid_positions_covers_every_wheel_not_just_four():
    """The old hexbin layout zipped a fixed 4-slot grid against however many
    wheels existed -- zip() silently drops anything past the 4th, which a
    6-wheel truck (3-axle) would hit."""
    from pptx.util import Emu
    cells = _grid_positions(6, Emu(0), Emu(0), Emu(9000000), Emu(6000000))
    assert len(cells) == 6


def test_grid_positions_empty_for_zero_wheels():
    from pptx.util import Emu
    assert _grid_positions(0, Emu(0), Emu(0), Emu(1000000), Emu(1000000)) == []


# -------------------------------------------------------- output directory

def test_report_saved_inside_run_output_dir(tmp_path):
    cfg = _fr_rr_config(tmp_path)
    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())
    assert out.exists()
    assert out.parent == cfg.run_output_dir.resolve()
    assert out.suffix == ".pptx"


# ------------------------------------------------------- no blank slides

def test_no_title_only_blank_slides_with_missing_figures_and_partial_wheels(tmp_path):
    """The originally-reported symptom, reproduced directly: FR/RR-only
    data with no figure files on disk at all (worst case) must still leave
    every slide with either real content or an explanatory message."""
    cfg = _fr_rr_config(tmp_path)
    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())
    prs = Presentation(str(out))
    assert _audit_slides(prs) == []


# ------------------------------------------------------- Apollo template

def test_report_uses_the_apollo_template_layouts(tmp_path):
    """Cover, dividers and content slides come from the charter deck's own
    layouts, 16:9, with no empty "Click to add" placeholders left behind."""
    from dtt.reporting.report_builder import (TEMPLATE, _LAYOUT_CONTENT, _LAYOUT_COVER,
                                              _LAYOUT_DIVIDER, _LAYOUT_RESULT,
                                              _LAYOUT_SECTION)
    assert TEMPLATE.exists()
    cfg = _fr_rr_config(tmp_path)
    out = build_report(cfg, _fr_rr_validation_report(), _fr_rr_stats(), _fr_rr_metadata())
    prs = Presentation(str(out))
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[0] == _LAYOUT_COVER
    assert layouts.count(_LAYOUT_DIVIDER) == 2          # "Data & Processing", "Module 1"
    assert layouts.count(_LAYOUT_SECTION) == 1          # "PV Format"
    # the charter-format result slides: two PV pages, box plots, combination
    assert layouts.count(_LAYOUT_RESULT) >= 4
    assert layouts.count(_LAYOUT_CONTENT) == (len(layouts) - 4
                                              - layouts.count(_LAYOUT_RESULT))
    for i, slide in enumerate(prs.slides, start=1):
        for ph in slide.placeholders:
            assert ph.has_text_frame and ph.text_frame.text.strip(), (i, ph.name)
